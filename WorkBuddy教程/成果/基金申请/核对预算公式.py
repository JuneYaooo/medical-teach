# -*- coding: utf-8 -*-
"""预算表公式独立核对：纯 Python 求值器，逐格重算并与缓存值比对"""
import math, re
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.utils import column_index_from_string

from pathlib import Path
PATH = str(Path(__file__).with_name('预算及测算依据.xlsx'))
wbf = load_workbook(PATH, data_only=False)
wbv = load_workbook(PATH, data_only=True)

def colrow(ref):
    c, r = coordinate_from_string(ref)
    return column_index_from_string(c), r

def excel_round(x, n=0):
    x = float(x); n = int(n)
    return math.floor(abs(x) * 10**n + 0.5) / 10**n * (1 if x >= 0 else -1)

def if_(cond, a, b):
    return a if cond else b

def or_(*a):
    return any(bool(x) for x in a)

SUMIF_PAT = re.compile(
    r"SUMIF\((?:(\w+)!)?([A-Z]+)(\d+):[A-Z]+(\d+),([^,]+),(?:(\w+)!)?([A-Z]+)(\d+):[A-Z]+(\d+)\)")

def eval_formula(sheet, formula):
    f = formula.lstrip('=')
    f = re.sub(r"'([^']+)'!", r'\1!', f)   # 规范化跨表前缀
    f = f.replace('$', '')                  # 去绝对引用符

    def sumif(m):
        s1, colA, r1, r2, crit, s2, colC, r3, r4 = m.groups()
        s1 = s1 or sheet; s2 = s2 or sheet
        crit = crit.strip().strip('"')
        ca = column_index_from_string(colA)
        cc = column_index_from_string(colC)
        tot = 0
        for r in range(int(r1), int(r2) + 1):
            v = wbv[s1].cell(row=r, column=ca).value
            if str(v) == crit:
                sv = wbv[s2].cell(row=r, column=cc).value
                if isinstance(sv, (int, float)):
                    tot += sv
        return repr(float(tot))

    f = SUMIF_PAT.sub(sumif, f)

    def sumr(m):
        s1, a, b = m.group(1) or sheet, m.group(2), m.group(3)
        ca, ra = colrow(a); cb, rb = colrow(b)
        tot = 0
        for r in range(ra, rb + 1):
            for c in range(ca, cb + 1):
                v = wbv[s1].cell(row=r, column=c).value
                if isinstance(v, (int, float)):
                    tot += v
        return repr(float(tot))

    f = re.sub(r'SUM\((?:(\w+)!)?([A-Z]+\d+):([A-Z]+\d+)\)', sumr, f)

    def ref(m):
        s1, col, row = m.group(1) or sheet, m.group(2), int(m.group(3))
        v = wbv[s1].cell(row=row, column=column_index_from_string(col)).value
        if v is None:
            return '0'
        if isinstance(v, (int, float)):
            return repr(float(v))
        return repr(str(v))

    f = re.sub(r'(?:(\w+)!)?([A-Z]{1,2})(\d+)\b', ref, f)
    f = f.replace('""', "''").replace('<>', '!=')
    f = re.sub(r'(?<![<>=!])=(?!=)', '==', f)
    f = f.replace('ROUND', 'excel_round').replace('IF(', 'if_(').replace('OR(', 'or_(')
    return eval(f)

def main():
    checked = mismatch = errs = 0
    rows = []
    for sheet in wbf.sheetnames:
        for row in wbf[sheet].iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    cached = wbv[sheet].cell(row=cell.row, column=cell.column).value
                    try:
                        calc = eval_formula(sheet, cell.value)
                    except Exception as e:
                        errs += 1
                        print('EVAL-ERR', sheet, cell.coordinate, cell.value, '|', e)
                        continue
                    checked += 1
                    ok = False
                    try:
                        if isinstance(cached, str):
                            ok = str(calc) == cached
                        else:
                            ok = cached is not None and abs(float(calc) - float(cached)) < 0.01
                    except Exception:
                        ok = False
                    if not ok:
                        mismatch += 1
                        print('MISMATCH', sheet, cell.coordinate, cell.value,
                              '| cached=', cached, '| calc=', calc)
                    rows.append((sheet, cell.coordinate, ok))
    print('公式总数:', checked, '| 独立重算一致:', checked - mismatch,
          '| 不一致:', mismatch, '| 求值失败:', errs)

if __name__ == '__main__':
    main()
