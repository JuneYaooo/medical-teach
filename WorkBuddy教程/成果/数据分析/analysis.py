# -*- coding: utf-8 -*-
"""
家庭血压项目 SAP 分析脚本（按《资料/数据字典与SAP.md》预设）
流程：清洗 -> clean.csv + 逐条清洗日志 -> Table1(基线) + 过程与随访描述 -> 完整病例线性回归
      （调整基线SBP/年龄/中心）-> HC3 敏感性 -> 模型诊断 -> 纳入 vs 未纳入比较 -> 次要结局描述
      -> 图与 Excel（含排版格式）。
推断说明：OLS 主分析使用 t 分布；HC3 敏感性使用正态近似（statsmodels cov_type='HC3' 默认 z 推断）。
说明：完整病例分析（complete-case），不称为完整 ITT；无预设亚组；不做均数填补。
运行：python3 analysis.py  （在 成果/数据分析/ 目录下生成全部交付物）
"""
import pandas as pd
import numpy as np
import statsmodels.formula.api as smf
from statsmodels.stats.diagnostic import het_breuschpagan
from scipy import stats
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pathlib import Path
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

plt.rcParams["font.sans-serif"] = ["Hiragino Sans GB", "PingFang SC", "Arial Unicode MS"]
plt.rcParams["axes.unicode_minus"] = False

BASE = Path(__file__).resolve().parent.parent.parent    # 项目根（脚本位于 成果/数据分析/）
OUT = BASE / "成果" / "数据分析"
OUT.mkdir(parents=True, exist_ok=True)
RAW = BASE / "资料" / "研究数据.csv"
SENTINELS = {"age": 999, "bmi": 99, "sbp_6m": 18}      # 字典/已确认清洗规则

# ---------- 1. 清洗 ----------
raw = pd.read_csv(RAW, encoding="utf-8-sig")
log = []  # 逐条清洗日志
def log_op(id_, line, field, old, new, rule, note=""):
    log.append({"编号": id_, "原始CSV数据行号": line, "字段": field,
                "原值": old, "新值": new, "规则": rule, "备注": note})

n_raw = len(raw)
dup_ids = raw["id"][raw["id"].duplicated()].unique().tolist()
drop_lines, conflict_pending = [], []
for pid in dup_ids:
    g = raw[raw["id"] == pid]
    identical = g.drop(columns=["id"]).nunique(dropna=False).eq(1).all()
    if identical:
        keep_line = g.index[0] + 1                      # 数据行号 = 索引+1
        for idx in g.index[1:]:
            drop_lines.append(idx)
            log_op(pid, idx + 1, "整行", "完全重复副本", "删除",
                   "规则C：仅移除逐字段完全一致的重复记录", f"与第{keep_line}行15个字段完全一致")
    else:
        conflict_pending.append(pid)                    # 冲突记录单列待核，不移除
assert n_raw - len(drop_lines) == 300, "清洗后应剩300行"
df = raw.drop(index=drop_lines).copy().reset_index(drop=True)

sex_map = {"男": 1, "M": 1, "1": 1, "女": 0, "F": 0, "0": 0}
for idx, v in df["sex"].items():
    if str(v) not in sex_map:
        log_op(df.at[idx, "id"], idx + 1, "sex", v, "待核",
               "规则B：性别按字典统一", "无法识别的取值，未改动，列入待核")
        conflict_pending.append(f"{df.at[idx,'id']}(sex={v})")
        continue
    if str(v) not in ("1", "0"):
        log_op(df.at[idx, "id"], idx + 1, "sex", v, sex_map[str(v)],
               "规则B：性别按字典统一为原列1/0（男/M/1→1，女/F/0→0）")
df["sex"] = df["sex"].map(lambda v: sex_map[str(v)])

for field, sv in SENTINELS.items():
    for idx in df.index[df[field] == sv]:
        log_op(df.at[idx, "id"], idx + 1, field, sv, "空(NaN)",
               f"规则A：{field}哨兵值{sv}设为缺失，保留受试者")
        df.at[idx, field] = np.nan

log_df = pd.DataFrame(log)
log_df.insert(0, "序号", range(1, len(log_df) + 1))
log_df.to_csv(OUT / "清洗执行日志.csv", index=False, encoding="utf-8-sig")
if conflict_pending:
    pd.DataFrame({"待核": conflict_pending}).to_csv(OUT / "冲突待核清单.csv",
                                                    index=False, encoding="utf-8-sig")
df.to_csv(OUT / "clean.csv", index=False, encoding="utf-8-sig")

assert len(df) == 300 and df["id"].nunique() == 300
assert df["sex"].isin([0, 1]).all() and not df["sex"].isna().any()
assert df["age"].max() < 999 and df["bmi"].max() < 99 and (df["sbp_6m"].dropna() != 18).all()
print(f"[清洗] 原始{n_raw}行 -> clean.csv {len(df)}行；日志{len(log_df)}条"
      f"（删重复{len(drop_lines)}，性别重编码{(log_df['字段']=='sex').sum()}，哨兵值{sum(log_df['字段'].isin(SENTINELS))}）；"
      f"待核：{conflict_pending if conflict_pending else '无'}")

# ---------- 2. 分析集、Table1（基线）与过程/随访描述 ----------
model_vars = ["sbp_6m", "group", "sbp_baseline", "age", "site"]
df["in_main"] = df[model_vars].notna().all(axis=1)
ana = df[df["in_main"]].copy()
n_in, n_out = int(df["in_main"].sum()), int((~df["in_main"]).sum())
print(f"[分析集] 完整病例 n={n_in}（未纳入 n={n_out}；完整病例分析，非完整ITT）")

T1_COLS = ["变量（单位）", "统计口径", "干预组（n=150）", "对照组（n=150）", "全体（N=300）", "缺失说明"]

def cont_cell(d, col, dec=1):
    v = d[col].dropna()
    return f"{v.mean():.{dec}f} ± {v.std():.{dec}f}（n={len(v)}）" if len(v) else "—"

def cat_cell(d, col, val, label):
    v = d[col].dropna(); n = len(v)
    cnt = int((v == val).sum())
    return f"{label} {cnt} ({cnt/n*100:.1f}%)" if n else f"{label} {cnt}"

def build_rows(variables):
    rows = []
    for spec in variables:
        col, unit, kind, dec, val, lab = spec
        if kind == "cont":
            name = f"{col}（{unit}）"
            stat = "均值 ± SD（有效n）"
            cells = [cont_cell(df[df.group=="干预组"], col, dec),
                     cont_cell(df[df.group=="对照组"], col, dec),
                     cont_cell(df, col, dec)]
            miss = f"缺失{int(df[col].isna().sum())}"
        else:
            name = f"{col}（{unit}）" if unit else col
            stat = "n (%)"
            cells = [cat_cell(df[df.group=="干预组"], col, val, lab),
                     cat_cell(df[df.group=="对照组"], col, val, lab),
                     cat_cell(df, col, val, lab)]
            miss = "无缺失" if df[col].isna().sum() == 0 else f"缺失{int(df[col].isna().sum())}"
        rows.append([name, stat] + cells + [miss])
    return pd.DataFrame(rows, columns=T1_COLS)

baseline_vars = [
    ("age", "岁", "cont", 1, None, None),
    ("bmi", "kg/m²", "cont", 1, None, None),
    ("duration_years", "年", "cont", 1, None, None),
    ("sbp_baseline", "mmHg", "cont", 1, None, None),
    ("dbp_baseline", "mmHg", "cont", 1, None, None),
    ("sex", "", "cat", 1, 1, "男"),
    ("diabetes", "", "cat", 1, 1, "有"),
    ("site", "", "cat", 1, "B", "B中心"),
]
table1 = build_rows(baseline_vars)
table1.to_csv(OUT / "Table1.csv", index=False, encoding="utf-8-sig")

process_vars = [
    ("adherence", "0–1，记录依从率", "cont", 2, None, None),
    ("visits", "次", "cont", 1, None, None),
]
process_tbl = build_rows(process_vars)
# 6个月结局缺失口径（清洗后，分母=各组150/全体300）：
#   sbp_6m 缺失 31 人 = 原始三者同步缺失 30 人 + SIM098（哨兵值设缺失）；
#   SIM098 清洗后仅 sbp_6m 与 controlled_6m 缺失，dbp_6m=72.7 仍保留，故 dbp_6m 缺失 30 人。
n6i, n6c = int(df[df.group=="干预组"]["sbp_6m"].isna().sum()), int(df[df.group=="对照组"]["sbp_6m"].isna().sum())
ndbi, ndbc = int(df[df.group=="干预组"]["dbp_6m"].isna().sum()), int(df[df.group=="对照组"]["dbp_6m"].isna().sum())
nci, ncc = int(df[df.group=="干预组"]["controlled_6m"].isna().sum()), int(df[df.group=="对照组"]["controlled_6m"].isna().sum())
assert (n6i+n6c, ndbi+ndbc, nci+ncc) == (31, 30, 31), "6个月缺失口径应为 sbp_6m=31, dbp_6m=30, controlled_6m=31"
def miss_cell(i, tot=150):
    return f"缺失 {i} ({i/tot*100:.1f}%)"
process_tbl.loc[len(process_tbl)] = ["sbp_6m（6个月收缩压，mmHg）缺失", "n (%)",
    miss_cell(n6i), miss_cell(n6c), miss_cell(n6i+n6c, 300),
    "31人=原始sbp_6m/dbp_6m/controlled_6m三者同步缺失30人 + SIM098（哨兵值18设缺失）"]
process_tbl.loc[len(process_tbl)] = ["dbp_6m（6个月舒张压，mmHg）缺失", "n (%)",
    miss_cell(ndbi), miss_cell(ndbc), miss_cell(ndbi+ndbc, 300),
    "30人；SIM098的dbp_6m=72.7仍保留，未缺失"]
process_tbl.loc[len(process_tbl)] = ["controlled_6m（6个月达标，<140/90）缺失", "n (%)",
    miss_cell(nci), miss_cell(ncc), miss_cell(nci+ncc, 300),
    "31人 = 上述三者同步缺失30人 + SIM098（sbp_6m设缺失后无法判定达标）"]
process_tbl.to_csv(OUT / "过程与随访描述.csv", index=False, encoding="utf-8-sig")
print("[Table1] 基线特征已生成（分母=清洗后300例）；过程与随访描述单独成表"
      "（adherence/visits/6个月三变量缺失：sbp_6m=31、dbp_6m=30、controlled_6m=31）")

# ---------- 3. 主分析：完整病例线性回归 ----------
fml = 'sbp_6m ~ C(group, Treatment(reference="对照组")) + sbp_baseline + age + C(site)'
fit = smf.ols(fml, data=ana).fit()
fit_hc3 = smf.ols(fml, data=ana).fit(cov_type="HC3")
gt = [k for k in fit.params.index if "干预组" in k][0]

def ci_row(res, name):
    b = res.params[name]
    lo, hi = res.conf_int().loc[name, 0], res.conf_int().loc[name, 1]
    return dict(系数=b, CI下限=lo, CI上限=hi, P值=res.pvalues[name])

main_res, hc3_res = ci_row(fit, gt), ci_row(fit_hc3, gt)

# 边际标准化（g-computation）：将全部268名完整病例分别整体设为干预/对照，
# 保留每例自身的年龄、基线SBP与中心，预测后取平均。
Xg = ana.copy(); Xg["group"] = "干预组"; adj_mean_int = float(fit.predict(Xg).mean())
Xg["group"] = "对照组"; adj_mean_ctl = float(fit.predict(Xg).mean())
adj_diff = adj_mean_int - adj_mean_ctl
print(f"[主分析] n={n_in}; 调整均数差(干预-对照)={main_res['系数']:.2f} mmHg "
      f"(95%CI {main_res['CI下限']:.2f} 至 {main_res['CI上限']:.2f}), P={main_res['P值']:.3g}; R2={fit.rsquared:.3f}")
print(f"[HC3]    β={hc3_res['系数']:.2f} (95%CI {hc3_res['CI下限']:.2f} 至 {hc3_res['CI上限']:.2f}), P={hc3_res['P值']:.3g}"
      "；推断：OLS=t分布，HC3=正态近似")
print(f"[边际标准化调整均数] 干预组 {adj_mean_int:.2f} vs 对照组 {adj_mean_ctl:.2f} mmHg（差 {adj_diff:.4f}，应与组系数一致）")

# 系数表：保留底层数值精度，CSV 用 8 位有效数字
coef_tbl = pd.DataFrame({"系数": fit.params, "SE(常规)": fit.bse,
                         "95%CI下限": fit.conf_int()[0], "95%CI上限": fit.conf_int()[1],
                         "P值": fit.pvalues, "SE(HC3)": fit_hc3.bse,
                         "P值(HC3)": fit_hc3.pvalues})
coef_tbl.index.name = "项"
coef_tbl.to_csv(OUT / "模型系数表.csv", encoding="utf-8-sig", float_format="%.8g")
assert (coef_tbl["P值"] > 0).all(), "P值不得为0"

# 模型诊断
bp_stat, bp_p, _, _ = het_breuschpagan(fit.resid, fit.model.exog)
sw_stat, sw_p = stats.shapiro(fit.resid)
resid_sd = float(fit.resid.std(ddof=1))          # 残差样本SD（268个残差，ddof=1）
rse = float(np.sqrt(fit.mse_resid))              # 模型残差标准误 RSE = sqrt(SSE/(n-k))，k=5，df=263
max_cook = float(fit.get_influence().cooks_distance[0].max())
print(f"[诊断] Breusch-Pagan P={bp_p:.4f}; Shapiro-Wilk P={sw_p:.4f}; 最大Cook's D={max_cook:.3f}")
print(f"[精度] 残差样本SD={resid_sd:.4f} mmHg（ddof=1）; 模型残差标准误RSE={rse:.4f} mmHg（sqrt(SSE/(n-k))，df={int(fit.df_resid)}）")

fig, ax = plt.subplots(2, 2, figsize=(11, 8.5))
ax[0,0].scatter(fit.fittedvalues, fit.resid, s=18, alpha=0.7, color="#3568a4")
ax[0,0].axhline(0, ls="--", c="r", lw=1)
ax[0,0].set(xlabel="拟合值 (mmHg)", ylabel="残差 (mmHg)", title="残差 vs 拟合值")
stats.probplot(fit.resid, dist="norm", plot=ax[0,1])
ax[0,1].set_title("残差 Q-Q 图")
s = np.sqrt(np.abs(fit.resid / fit.resid.std()))
ax[1,0].scatter(fit.fittedvalues, s, s=18, alpha=0.7, color="#3568a4")
ax[1,0].set(xlabel="拟合值 (mmHg)", ylabel="√|标准化残差|", title="尺度-位置图")
ax[1,1].hist(fit.resid, bins=20, color="#3568a4", alpha=0.8)
ax[1,1].set(xlabel="残差 (mmHg)", ylabel="频数", title="残差分布")
fig.suptitle("主分析模型诊断（完整病例 n=%d）" % n_in)
fig.tight_layout(); fig.savefig(OUT / "模型诊断图.png", dpi=150); plt.close(fig)

# ---------- 4. 纳入 vs 未纳入比较 ----------
def cmp_cont(col):
    a = df.loc[df.in_main, col].dropna(); b = df.loc[~df.in_main, col].dropna()
    t, p = stats.ttest_ind(a, b, equal_var=False)
    return f"{a.mean():.1f} ± {a.std():.1f}", f"{b.mean():.1f} ± {b.std():.1f}", f"{p:.3f}", f"缺失不计入，n={len(a)}/{len(b)}"
def cmp_cat(col, val):
    a = df.loc[df.in_main, col]; b = df.loc[~df.in_main, col]
    tab = pd.crosstab(df.in_main, df[col] == val)
    p = stats.chi2_contingency(tab)[1] if tab.shape == (2, 2) else np.nan
    return f"{(a==val).sum()} ({(a==val).mean()*100:.1f}%)", f"{(b==val).sum()} ({(b==val).mean()*100:.1f}%)", f"{p:.3f}", f"n={len(a)}/{len(b)}"
cmp_rows = [("n", "", f"{n_in}", f"{n_out}", "", "纳入/未纳入")]
for col, unit in [("age","岁"),("sbp_baseline","mmHg"),("dbp_baseline","mmHg"),("bmi","kg/m²"),
                  ("duration_years","年"),("adherence","0–1"),("visits","次")]:
    a,b,p,nn = cmp_cont(col); cmp_rows.append((f"{col}（{unit}）, 均值±SD", "", a, b, p, nn))
for col, val, lab in [("sex",1,"男性"),("diabetes",1,"糖尿病"),("group","干预组","干预组")]:
    a,b,p,nn = cmp_cat(col,val); cmp_rows.append((f"{lab}, n (%)", "", a, b, p, nn))
a,b,p,nn = cmp_cat("site","B"); cmp_rows.append(("中心B, n (%)", "", a, b, p, nn))
cmp_tbl = pd.DataFrame(cmp_rows, columns=["变量","单位","纳入者(n=%d)"%n_in,"未纳入者(n=%d)"%n_out,"P值","说明"])
cmp_tbl.to_csv(OUT / "纳入与未纳入比较.csv", index=False, encoding="utf-8-sig")
print("[比较] 纳入vs未纳入基线特征已生成（Welch t / 卡方，仅描述两组可比性）")

# ---------- 5. 次要结局（仅描述） ----------
sec_raw = df[df["controlled_6m"].notna()].groupby("group")["controlled_6m"].agg(["count", "sum", "mean"])
sec_tbl = sec_raw.copy()
sec_tbl.columns = ["有效n", "达标n", "达标比例"]
sec_tbl["有效n"] = sec_tbl["有效n"].astype(int)      # 人数按整数显示
sec_tbl["达标n"] = sec_tbl["达标n"].astype(int)
sec_tbl.to_csv(OUT / "次要结局_达标比例.csv", encoding="utf-8-sig")
print("[次要结局]\n", sec_tbl)

# ---------- 6. 主结果图（带单位与CI） ----------
fig, ax = plt.subplots(figsize=(8.2, 3.6))
labels = ["主分析（OLS）", "敏感性分析（HC3 稳健SE）"]
est = [main_res["系数"], hc3_res["系数"]]
lo  = [main_res["CI下限"], hc3_res["CI下限"]]
hi  = [main_res["CI上限"], hc3_res["CI上限"]]
y = [1.6, 0.8]
for i in range(2):
    ax.errorbar(est[i], y[i], xerr=[[est[i]-lo[i]], [hi[i]-est[i]]], fmt="s",
                color="#c0392b" if est[i] < 0 else "#2c7a2c", capsize=5, markersize=8, lw=2)
    ax.text(hi[i]+0.6, y[i], f"{est[i]:.2f} ({lo[i]:.2f}, {hi[i]:.2f})", va="center", fontsize=10)
ax.axvline(0, ls="--", c="gray", lw=1)
ax.set_yticks(y); ax.set_yticklabels(labels)
ax.set_xlabel("调整均数差：干预组-对照组（mmHg，95% CI）")
ax.set_xlim(min(lo)-2.5, max(hi)+4.5)
ax.set_title(f"6 个月收缩压调整组间差（完整病例 n={n_in}，调整基线SBP、年龄、中心）")
ax.spines[["top","right"]].set_visible(False)
fig.tight_layout(); fig.savefig(OUT / "主结果图.png", dpi=150); plt.close(fig)
print("[图] 主结果图.png / 模型诊断图.png 已生成")

# ---------- 7. 结果 Excel（含排版格式） ----------
main_sheet = pd.DataFrame([
    {"指标": "干预组调整均数(mmHg)", "估计值": adj_mean_int,
     "定义": "边际标准化：全部268例设为干预组，保留各自年龄/基线SBP/中心，预测后取平均"},
    {"指标": "对照组调整均数(mmHg)", "估计值": adj_mean_ctl,
     "定义": "边际标准化：全部268例设为对照组，同上"},
    {"指标": "调整均数差(干预-对照, mmHg)", "估计值": main_res["系数"],
     "95%CI下限": main_res["CI下限"], "95%CI上限": main_res["CI上限"], "P值": main_res["P值"],
     "定义": "回归组系数；OLS推断基于t分布"},
    {"指标": "R²", "估计值": fit.rsquared, "定义": "决定系数"},
    {"指标": "残差样本SD(mmHg)", "估计值": resid_sd,
     "定义": "268个残差的样本标准差（ddof=1），描述残差离散度"},
    {"指标": "模型残差标准误RSE(mmHg)", "估计值": rse,
     "定义": "sqrt(SSE/(n-k))，k=5，df=263；模型误差项标准差估计"},
])
coef_sheet = coef_tbl.reset_index()
coef_sheet["推断说明"] = np.where(coef_sheet["项"] == gt,
    "组系数：OLS基于t分布；HC3列基于正态近似", "")
hc3_sheet = pd.DataFrame([
    {"分析": "主分析(OLS, t分布)", "调整均数差(mmHg)": main_res["系数"], "95%CI下限": main_res["CI下限"],
     "95%CI上限": main_res["CI上限"], "P值": main_res["P值"]},
    {"分析": "敏感性(HC3稳健SE, 正态近似)", "调整均数差(mmHg)": hc3_res["系数"], "95%CI下限": hc3_res["CI下限"],
     "95%CI上限": hc3_res["CI上限"], "P值": hc3_res["P值"]},
])
diag_sheet = pd.DataFrame([
    {"诊断": "Breusch-Pagan异方差检验", "统计量": bp_stat, "P值": bp_p},
    {"诊断": "Shapiro-Wilk正态性(残差)", "统计量": sw_stat, "P值": sw_p},
    {"诊断": "最大Cook距离", "统计量": max_cook, "P值": np.nan},
])

def disp_w(s):
    return sum(2 if ord(ch) > 127 else 1 for ch in str(s))

def format_sheet(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    for col_idx in range(1, ws.max_column + 1):
        header = str(ws.cell(1, col_idx).value or "")
        hl = header.lower()
        w = disp_w(header) + 3
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row, col_idx)
            v = cell.value
            if v is None:
                continue
            w = max(w, disp_w(v) + 2)
            if isinstance(v, str) and disp_w(v) > 22:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                if hl.startswith("p") or "p值" in hl:
                    cell.number_format = "0.00E+00"      # 科学计数法，P不为0
                elif "比例" in header:
                    cell.number_format = "0.0%"
                elif isinstance(v, float):
                    cell.number_format = "0.000"
        ws.column_dimensions[get_column_letter(col_idx)].width = min(58, max(10, w))

with pd.ExcelWriter(OUT / "结果汇总.xlsx", engine="openpyxl") as w:
    pd.DataFrame({
        "项目": ["原始记录数", "独立受试者", "清洗后行数", "完整病例分析n", "未纳入n",
                 "分析类型", "模型", "调整协变量", "组别参照", "调整均数定义",
                 "OLS推断分布", "HC3推断分布"],
        "值": [n_raw, 300, len(df), n_in, n_out, "完整病例分析（complete-case），不称为完整ITT",
               "线性回归", "基线收缩压、年龄、中心", "对照组（均数差=干预-对照）",
               "边际标准化（g-computation）：整体替换group后按各例自身协变量预测取平均",
               "t 分布", "正态近似（z）"]
    }).to_excel(w, sheet_name="说明", index=False)
    table1.to_excel(w, sheet_name="Table1_基线特征", index=False)
    process_tbl.to_excel(w, sheet_name="过程与随访", index=False)
    main_sheet.to_excel(w, sheet_name="主分析", index=False)
    coef_sheet.to_excel(w, sheet_name="模型系数", index=False)
    hc3_sheet.to_excel(w, sheet_name="HC3敏感性", index=False)
    cmp_tbl.to_excel(w, sheet_name="纳入vs未纳入", index=False)
    sec_tbl.reset_index().to_excel(w, sheet_name="次要结局", index=False)
    diag_sheet.to_excel(w, sheet_name="模型诊断", index=False)
    log_df.to_excel(w, sheet_name="清洗执行日志", index=False)
    for ws in w.book.worksheets:
        format_sheet(ws)
print("[Excel] 结果汇总.xlsx 已生成（10个sheet，含列宽/冻结表头/科学计数法P值/数值格式）")

# ---------- 8. 复核输出 ----------
print(f"[复核] clean.csv 行数={len(pd.read_csv(OUT/'clean.csv'))}; 日志条数={len(log_df)}; "
      f"主n={n_in}; MD={main_res['系数']:.4f} [{main_res['CI下限']:.4f},{main_res['CI上限']:.4f}] "
      f"P={main_res['P值']:.4e}; HC3 MD={hc3_res['系数']:.4f} [{hc3_res['CI下限']:.4f},{hc3_res['CI上限']:.4f}] "
      f"P={hc3_res['P值']:.4e}")
print(f"[复核] 边际标准化调整均数: 干预={adj_mean_int:.4f}, 对照={adj_mean_ctl:.4f}, 差={adj_diff:.4f}")
print("[复核] Table1干预组n=%d, 对照组n=%d; 次要结局达标: %s" % (
    int((df.group=='干预组').sum()), int((df.group=='对照组').sum()),
    "; ".join(f"{g} {int(r['sum'])}/{int(r['count'])} ({r['mean']*100:.1f}%)" for g, r in sec_raw.iterrows())))
