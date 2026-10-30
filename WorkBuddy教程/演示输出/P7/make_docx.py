#!/usr/bin/env python3
"""P7 患教单 DOCX 生成：A4、大字、仅事实卡 F01-F10 内容、填写栏、教学草稿标识。"""
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

from pathlib import Path
OUT = Path(__file__).resolve().parent / "P7-患教单.docx"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin = sec.bottom_margin = Mm(10)
sec.left_margin = sec.right_margin = Mm(11)

style = doc.styles["Normal"]
style.font.name = "PingFang SC"
style.font.size = Pt(12.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
pf = style.paragraph_format
pf.space_after = Pt(2)
pf.line_spacing = 1.15

def para(text="", size=12.5, bold=False, color=None, align=None, space_after=2, indent=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.left_indent = Pt(indent)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor(*color)
    return p

def heading(text):
    para(text, size=14, bold=True, space_after=1)

def bullet(text, bold_part=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1.5)
    if bold_part:
        i = text.find(bold_part)
        r1 = p.add_run(text[:i]); r2 = p.add_run(bold_part); r3 = p.add_run(text[i+len(bold_part):])
        r2.bold = True
        for r in (r1, r2, r3):
            r.font.size = Pt(12.5)
    else:
        r = p.add_run(text); r.font.size = Pt(12.5)

# ── 页眉标题 ──
para("家庭血压监测 · 患教单", size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
para("编号 P7 ｜ 核对日期 2026-10-30 ｜ 核对状态：教学来源核对", size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3, color=(68, 68, 68))
para("教学草稿 —— 医生签署：未签署。仅供教学演示，不可直接发给患者使用。",
     size=11.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6, color=(179, 0, 0))

# ── 一 ──
heading("一、在家测血压，不能代替看医生")
bullet("家庭监测不能替代医生随访，请按医生安排定期复诊。（F01）", bold_part="不能替代")

# ── 二 ──
heading("二、测量前 30 分钟")
bullet("不吸烟、不喝含咖啡因的饮料（如咖啡、浓茶）、不运动。（F04）")
bullet("测前排空膀胱，坐下来安静休息至少 5 分钟再测。（F05）", bold_part="至少 5 分钟")

# ── 三 ──
heading("三、怎么测")
bullet("用经验证的上臂式自动血压计，袖带大小要合适。（F03）", bold_part="经验证的上臂式")
bullet("袖带直接贴在上臂皮肤上，不要隔着衣服测。（F06）", bold_part="不要隔着衣服")
bullet("手臂要有支撑（如放在桌上），袖带中部与心脏同一高度。（F07）", bold_part="同一高度")
bullet("测量时保持安静：不说话、不用手机。（F08）", bold_part="不说话、不用手机")
bullet("每次测量取两次读数，两次之间隔 1 分钟，并把结果记下来。（F09）",
       bold_part="每次测量取两次读数，两次之间隔 1 分钟")

# ── 四 ──
heading("四、记下来，带去复诊")
bullet("把测量记录带上（最好连血压计一起），和医疗人员一起核对。（F10）")
bullet("不要仅因家里测的读数变化，就自行停药；任何用药调整，先和医疗人员商量。（F02）",
       bold_part="不要仅因家里测的读数变化，就自行停药")

# ── 五 填写栏 ──
heading("五、请医生填写（我还没有医嘱，请勿自行填写）")
rows = [
    ("我的随访日期", "＿＿＿＿＿＿年＿＿月＿＿日"),
    ("科室电话", "＿＿＿＿＿＿＿＿＿＿＿＿＿＿"),
    ("我一天测几次、什么时间测\n（具体频次由医疗人员确定）", "＿＿＿＿＿＿＿＿＿＿＿＿＿＿"),
    ("医生签名 / 盖章", "＿＿＿＿＿＿＿＿（未签署）"),
]
tbl = doc.add_table(rows=len(rows), cols=2)
tbl.style = "Table Grid"
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl.columns[0].width = Mm(62)
tbl.columns[1].width = Mm(115)
for i, (k, v) in enumerate(rows):
    c0, c1 = tbl.rows[i].cells
    c0.width, c1.width = Mm(62), Mm(115)
    for cell, text, size in ((c0, k, 12.5), (c1, v, 12.5)):
        cp = cell.paragraphs[0]
        for j, line in enumerate(text.split("\n")):
            if j > 0:
                cp = cell.add_paragraph()
            r = cp.add_run(line)
            r.font.size = Pt(size if j == 0 else 8.5)
            if j > 0:
                r.font.color.rgb = RGBColor(68, 68, 68)

# ── 来源脚注 ──
para("", space_after=1)
src = doc.add_paragraph()
src.paragraph_format.space_before = Pt(4)
r = src.add_run("来源：美国心脏协会（AHA）\u201cMonitoring your blood pressure at home\u201d（home monitoring / how to use a home blood pressure monitor），"
                "教学来源核对 2026-10-30。本单只包含该来源中的内容：不含急诊/就医阈值，不含停药建议；个体情况（包括一天测几次）请遵医嘱。"
                "修订说明：原稿\u201c测两次\u201d经患者回述发现易被误解为\u201c一天测两次\u201d，已改为\u201c每次测量取两次读数，间隔 1 分钟；一天测几次按医生安排\u201d。")
r.font.size = Pt(8)
r.font.color.rgb = RGBColor(68, 68, 68)

doc.save(OUT)
print("saved:", OUT)
